import io
from collections.abc import Sequence

from docx import Document

from app.config import Settings
from app.models import (
    ClauseStatus,
    Evidence,
    Finding,
    GroundingStatus,
    ReviewReport,
    RuleFinding,
    Severity,
)
from app.parsers import DocumentParser
from app.sample import sample_contract_bytes, sample_metadata
from app.service import ReviewService


class UntrustedSynthesis:
    def synthesize(
        self,
        rule_findings: Sequence[RuleFinding],
        evidence: Sequence[Evidence],
    ) -> tuple[str, list[Finding], str]:
        source = rule_findings[0]
        return (
            "Model-authored summary.",
            [
                Finding(
                    finding_id="F-001",
                    rule_id=source.rule_id,
                    title=source.title,
                    severity=source.severity,
                    description=source.description,
                    recommendation=source.recommendation,
                    citation_ids=[evidence[0].evidence_id, "EV-HALLUCINATED"],
                )
            ],
            "test-model",
        )


class UngroundedSynthesis:
    def __init__(self, citations: list[str]) -> None:
        self.citations = citations

    def synthesize(
        self,
        rule_findings: Sequence[RuleFinding],
        evidence: Sequence[Evidence],
    ) -> tuple[str, list[Finding], str]:
        source = rule_findings[0]
        return (
            "Untrusted summary.",
            [
                Finding(
                    finding_id="F-001",
                    rule_id=source.rule_id,
                    title=source.title,
                    severity=Severity.CRITICAL,
                    description=source.description,
                    recommendation=source.recommendation,
                    citation_ids=self.citations,
                )
            ],
            "test-model",
        )


def test_end_to_end_synthetic_review_is_offline_and_traceable() -> None:
    settings = Settings(bedrock_enabled=False)
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())

    report = ReviewService(settings).review(document, sample_metadata())

    assert report.synthesis_mode == "offline-deterministic"
    assert any(finding.severity == Severity.CRITICAL for finding in report.findings)
    assert any(clause.clause_id == "FAR 52.249" for clause in report.clause_inventory)
    assert all(clause.classifier_model_id for clause in report.clause_inventory)
    assert "local-heuristic-v1" in report.classifier_model_ids
    assert report.ablation_summary.classifier_assisted_clause_count >= (
        report.ablation_summary.deterministic_clause_count
    )
    assert report.evidence
    valid_ids = {item.evidence_id for item in report.evidence}
    assert all(set(finding.citation_ids) <= valid_ids for finding in report.findings)
    assert any(finding.locations for finding in report.findings)
    finding_by_rule = {finding.rule_id: finding for finding in report.findings}
    assert "EV-FAR-RISK" in finding_by_rule["GC-005"].citation_ids
    assert "EV-FAR-32" in finding_by_rule["GC-006"].citation_ids


def test_invalid_model_citation_ids_are_removed() -> None:
    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())

    report = ReviewService(settings, synthesis=UntrustedSynthesis()).review(
        document, sample_metadata()
    )

    assert report.findings[0].citation_ids
    assert "EV-HALLUCINATED" not in report.findings[0].citation_ids
    assert report.findings[0].grounding_status == GroundingStatus.VERIFIED


def test_globally_real_but_rule_unrelated_citation_is_rejected() -> None:
    class TwoRules:
        def evaluate(self, _document, _metadata, _clauses):
            return [
                RuleFinding(
                    rule_id="R-ONE",
                    title="First",
                    severity=Severity.HIGH,
                    description="First issue.",
                    recommendation="Review first.",
                    evidence_queries=["first"],
                ),
                RuleFinding(
                    rule_id="R-TWO",
                    title="Second",
                    severity=Severity.HIGH,
                    description="Second issue.",
                    recommendation="Review second.",
                    evidence_queries=["second"],
                ),
            ]

    class ScopedRetrieval:
        def retrieve(self, queries):
            suffix = queries[0].upper()
            return [
                Evidence(
                    evidence_id=f"EV-{suffix}",
                    source="Official",
                    title=suffix,
                    excerpt=f"Evidence for {suffix}.",
                )
            ]

    class CrossCitingSynthesis:
        def synthesize(self, _rules, _evidence):
            return (
                "Cross-cited.",
                [
                    Finding(
                        finding_id="F-001",
                        rule_id="R-ONE",
                        title="First",
                        severity=Severity.HIGH,
                        description="First issue.",
                        recommendation="Review first.",
                        citation_ids=["EV-SECOND"],
                    )
                ],
                "test-model",
            )

    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())
    report = ReviewService(
        settings,
        rules=TwoRules(),
        retrieval=ScopedRetrieval(),
        synthesis=CrossCitingSynthesis(),
    ).review(document, sample_metadata())

    assert {item.evidence_id for item in report.evidence} == {
        "EV-FIRST",
        "EV-SECOND",
    }
    assert report.findings[0].citation_ids == []
    assert report.findings[0].grounding_status == GroundingStatus.UNVERIFIED
    assert report.findings[0].severity == Severity.INFO


def test_hallucinated_only_citations_are_downgraded_and_labeled() -> None:
    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())
    report = ReviewService(
        settings,
        synthesis=UngroundedSynthesis(["EV-HALLUCINATED"]),
    ).review(document, sample_metadata())

    finding = report.findings[0]
    assert finding.citation_ids == []
    assert finding.severity == Severity.INFO
    assert finding.grounding_status == GroundingStatus.UNVERIFIED
    assert finding.description.startswith("UNVERIFIED —")


def test_empty_citations_are_downgraded_and_labeled() -> None:
    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())
    report = ReviewService(
        settings,
        synthesis=UngroundedSynthesis([]),
    ).review(document, sample_metadata())

    assert report.findings[0].severity == Severity.INFO
    assert report.findings[0].grounding_status == GroundingStatus.UNVERIFIED


def test_sample_clause_status_inventory_is_truthful() -> None:
    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())
    report = ReviewService(settings).review(document, sample_metadata())
    statuses = {item.clause_id: item.status for item in report.clause_status_inventory}

    assert statuses["FAR 52.249"] == ClauseStatus.PRESENT
    assert statuses["FAR 52.204-7"] == ClauseStatus.MISSING
    assert statuses["FAR 52.212-4"] == ClauseStatus.NOT_APPLICABLE
    present = next(
        item for item in report.clause_status_inventory if item.clause_id == "FAR 52.249"
    )
    assert present.detected_locations
    assert present.confidence is not None
    assert present.classifier_model_ids


def test_cots_only_does_not_trigger_dfars_or_generic_clause_mandates() -> None:
    contract = Document()
    contract.add_paragraph("Statement of work for standard catalog supplies.")
    buffer = io.BytesIO()
    contract.save(buffer)
    settings = Settings()
    document = DocumentParser(settings).parse("cots.docx", buffer.getvalue())
    metadata = sample_metadata().model_copy(update={"commercial_product": True, "cots_only": True})

    report = ReviewService(settings).review(document, metadata)

    assert not {"GC-001", "GC-002", "GC-003", "GC-004"} & {
        finding.rule_id for finding in report.findings
    }
    cyber = next(
        item for item in report.clause_status_inventory if item.clause_id == "DFARS 252.204-7012"
    )
    assert cyber.status == ClauseStatus.NOT_APPLICABLE


def test_uncertain_dfars_conditions_remain_undetermined() -> None:
    contract = Document()
    contract.add_paragraph("Statement of work.")
    buffer = io.BytesIO()
    contract.save(buffer)
    settings = Settings()
    document = DocumentParser(settings).parse("civilian.docx", buffer.getvalue())
    metadata = sample_metadata().model_copy(update={"agency": "Civilian Agency"})

    report = ReviewService(settings).review(document, metadata)

    cyber = next(
        item for item in report.clause_status_inventory if item.clause_id == "DFARS 252.204-7012"
    )
    assert cyber.status == ClauseStatus.UNDETERMINED


def test_review_report_serialization_round_trip() -> None:
    settings = Settings()
    document = DocumentParser(settings).parse("sample.docx", sample_contract_bytes())
    report = ReviewService(settings).review(document, sample_metadata())

    restored = ReviewReport.model_validate_json(report.model_dump_json())

    assert restored == report
    assert restored.clause_status_inventory


def test_malformed_clause_reference_is_reported_in_inventory() -> None:
    contract = Document()
    contract.add_paragraph("The contract includes DFARS 252.204.7012A.")
    buffer = io.BytesIO()
    contract.save(buffer)
    settings = Settings()
    document = DocumentParser(settings).parse("malformed.docx", buffer.getvalue())

    report = ReviewService(settings).review(document, sample_metadata())

    malformed = [
        item for item in report.clause_status_inventory if item.status == ClauseStatus.MALFORMED
    ]
    assert malformed
    assert malformed[0].detected_locations[0].label == "Page 1, paragraph 1"

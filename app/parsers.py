"""Safe in-memory PDF and DOCX extraction with reviewable locations."""

from __future__ import annotations

import hashlib
import io
import multiprocessing
import re
import sys
import zipfile
from collections.abc import Callable
from contextlib import suppress
from pathlib import Path, PurePosixPath
from typing import Any

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader

from app.config import Settings
from app.models import DocumentLocation, ExtractedSegment, ParsedDocument


def _apply_address_space_limit(memory_limit_mb: int) -> None:
    # macOS does not support lowering RLIMIT_AS reliably. The subprocess timeout
    # and explicit archive/page/text bounds remain active; Linux enforces memory.
    if __import__("os").name != "posix" or sys.platform == "darwin":
        return
    import resource

    requested = memory_limit_mb * 1024 * 1024
    _soft, hard = resource.getrlimit(resource.RLIMIT_AS)
    limit = requested if hard == resource.RLIM_INFINITY else min(requested, hard)
    resource.setrlimit(resource.RLIMIT_AS, (limit, limit))


def _parse_worker(
    filename: str,
    data: bytes,
    settings: Settings,
    connection: Any,
) -> None:
    try:
        _apply_address_space_limit(settings.parse_memory_limit_mb)
        parsed = DocumentParser(settings).parse(filename, data)
        connection.send(("ok", parsed.model_dump(mode="json")))
    except BaseException:
        # Child failures are intentionally opaque: parser exceptions can contain input.
        with suppress(BrokenPipeError, EOFError, OSError):
            connection.send(("error", None))
    finally:
        connection.close()


class DocumentParser:
    def __init__(
        self,
        settings: Settings,
        isolation_runner: Callable[[str, bytes], ParsedDocument] | None = None,
        worker_target: Callable[..., None] = _parse_worker,
    ) -> None:
        self.settings = settings
        self._isolation_runner = isolation_runner
        self._worker_target = worker_target

    async def read_upload(self, upload: UploadFile) -> bytes:
        extension = Path(upload.filename or "").suffix.lower()
        if extension not in self.settings.allowed_extensions:
            raise HTTPException(
                status_code=415,
                detail=f"Only {', '.join(self.settings.allowed_extensions)} files are allowed",
            )
        data = await upload.read(self.settings.max_upload_bytes + 1)
        await upload.close()
        if len(data) > self.settings.max_upload_bytes:
            raise HTTPException(status_code=413, detail="Uploaded file is too large")
        if not data:
            raise HTTPException(status_code=400, detail="Uploaded file is empty")
        return data

    def parse(self, filename: str, data: bytes) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        if extension not in self.settings.allowed_extensions:
            raise ValueError("Unsupported document type")
        try:
            segments = self._parse_pdf(data) if extension == ".pdf" else self._parse_docx(data)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Unable to parse {extension} document") from exc
        if not segments:
            raise ValueError("Document contains no extractable text")
        return ParsedDocument(
            filename=Path(filename).name,
            media_type=(
                "application/pdf"
                if extension == ".pdf"
                else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            sha256=hashlib.sha256(data).hexdigest(),
            segments=segments,
        )

    def parse_isolated(self, filename: str, data: bytes) -> ParsedDocument:
        if self._isolation_runner is not None:
            return self._isolation_runner(filename, data)
        context = multiprocessing.get_context("spawn")
        receive_connection, send_connection = context.Pipe(duplex=False)
        process = context.Process(
            target=self._worker_target,
            args=(filename, data, self.settings, send_connection),
            daemon=True,
        )
        started = False
        try:
            process.start()
            started = True
            send_connection.close()
            if not receive_connection.poll(self.settings.parse_timeout_seconds):
                self._stop_process(process)
                raise ValueError("Document parsing timed out")
            try:
                status, payload = receive_connection.recv()
            except (EOFError, OSError) as exc:
                raise ValueError("Document parsing failed in isolated worker") from exc
            process.join(timeout=1)
            if process.is_alive():
                self._stop_process(process)
            if status != "ok" or not isinstance(payload, dict):
                raise ValueError("Document parsing failed in isolated worker")
            return ParsedDocument.model_validate(payload)
        finally:
            receive_connection.close()
            send_connection.close()
            if started and process.is_alive():
                self._stop_process(process)

    @staticmethod
    def _stop_process(process: multiprocessing.Process) -> None:
        process.terminate()
        process.join(timeout=1)
        if process.is_alive() and hasattr(process, "kill"):
            process.kill()
            process.join(timeout=1)

    def _parse_pdf(self, data: bytes) -> list[ExtractedSegment]:
        segments: list[ExtractedSegment] = []
        extracted_characters = 0
        pages = PdfReader(io.BytesIO(data)).pages
        if len(pages) > self.settings.max_pdf_pages:
            raise ValueError("PDF page limit exceeded")
        for page_number, page in enumerate(pages, start=1):
            text = page.extract_text() or ""
            extracted_characters += len(text)
            if extracted_characters > self.settings.max_extracted_characters:
                raise ValueError("Extracted text limit exceeded")
            paragraphs = [
                part.strip() for part in re.split(r"\n\s*\n|(?<=\.)\s*\n", text) if part.strip()
            ]
            for paragraph_number, text in enumerate(paragraphs, start=1):
                if len(segments) >= self.settings.max_document_segments:
                    raise ValueError("Document segment limit exceeded")
                segment_id = f"p{page_number}-para{paragraph_number}"
                segments.append(
                    ExtractedSegment(
                        segment_id=segment_id,
                        text=text,
                        location=DocumentLocation(
                            page=page_number,
                            paragraph=paragraph_number,
                            label=f"Page {page_number}, paragraph {paragraph_number}",
                        ),
                    )
                )
        return segments

    def _parse_docx(self, data: bytes) -> list[ExtractedSegment]:
        self._preflight_docx(data)
        segments: list[ExtractedSegment] = []
        extracted_characters = 0
        page = 1
        paragraph_number = 0
        for paragraph in Document(io.BytesIO(data)).paragraphs:
            text = paragraph.text.strip()
            if not text:
                if 'w:type="page"' in paragraph._p.xml:
                    page += 1
                    paragraph_number = 0
                continue
            extracted_characters += len(text)
            if extracted_characters > self.settings.max_extracted_characters:
                raise ValueError("Extracted text limit exceeded")
            if len(segments) >= self.settings.max_document_segments:
                raise ValueError("Document segment limit exceeded")
            paragraph_number += 1
            segment_id = f"p{page}-para{paragraph_number}"
            segments.append(
                ExtractedSegment(
                    segment_id=segment_id,
                    text=text,
                    location=DocumentLocation(
                        page=page,
                        paragraph=paragraph_number,
                        label=f"Page {page}, paragraph {paragraph_number}",
                    ),
                )
            )
            if 'w:type="page"' in paragraph._p.xml:
                page += 1
                paragraph_number = 0
        return segments

    def _preflight_docx(self, data: bytes) -> None:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > self.settings.max_docx_members:
                raise ValueError("DOCX archive member limit exceeded")
            total_uncompressed = 0
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts or member.flag_bits & 0x1:
                    raise ValueError("Suspicious DOCX archive")
                total_uncompressed += member.file_size
                if (
                    member.file_size > 0
                    and member.compress_size > 0
                    and member.file_size / member.compress_size
                    > self.settings.max_docx_compression_ratio
                ):
                    raise ValueError("Suspicious DOCX compression ratio")
            if total_uncompressed > self.settings.max_docx_uncompressed_mb * 1024 * 1024:
                raise ValueError("DOCX uncompressed size limit exceeded")

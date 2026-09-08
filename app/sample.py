"""In-memory sample contract used by the judge workflow."""

import io

from docx import Document

from app.models import AcquisitionMetadata, AcquisitionStage, ContractType


def sample_contract_bytes() -> bytes:
    document = Document()
    document.add_heading("SAMPLE DEFENSE SERVICES CONTRACT", level=1)
    document.add_paragraph(
        "Solicitation DEMO-2026-001 provides engineering support for a twelve-month period."
    )
    document.add_paragraph("FAR 52.233-1 Disputes is incorporated by reference.")
    document.add_paragraph("Payment will be made 90 days after receipt of an accepted invoice.")
    document.add_paragraph(
        "The contractor accepts unlimited liability for all consequential damages."
    )
    document.add_paragraph("The Government may terminate this contract under FAR 52.249-2.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def sample_metadata() -> AcquisitionMetadata:
    return AcquisitionMetadata(
        agency="Department of Defense",
        solicitation_number="DEMO-2026-001",
        contract_type=ContractType.FIRM_FIXED_PRICE,
        estimated_value=1_500_000,
        set_aside="Small Business",
        commercial_product=False,
        cots_only=False,
        performance_months=12,
        place_of_performance="Arlington, Virginia",
        acquisition_stage=AcquisitionStage.SOLICITATION,
    )
